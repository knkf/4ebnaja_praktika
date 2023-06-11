import openpyxl
from PyQt5 import QtWidgets
from PyQt5.QtWidgets import QMainWindow, QApplication, QPushButton, QLineEdit, QMessageBox, QMenuBar, QMenu, QAction, \
    QTableWidget, QHeaderView, QTableWidgetItem, QLabel, QRadioButton, QTextEdit
from PyQt5 import uic, QtCore
import sys
import pandas as pd


class PersonalData(QMainWindow):
    def __init__(self, incoming_id, is_full):
        super(PersonalData, self).__init__()
        self.id = incoming_id
        self.is_full = is_full
        uic.loadUi('personalData.ui', self)

        self.passport_series = self.findChild(QLineEdit, 'passport_series')
        self.passport_id = self.findChild(QLineEdit, 'passport_id')
        self.birthday_place = self.findChild(QLineEdit, 'birthday_place')
        self.living_place = self.findChild(QLineEdit, 'living_place')
        self.family_status = self.findChild(QLineEdit, 'family_status')
        self.education = self.findChild(QLineEdit, 'education')
        self.work_place = self.findChild(QLineEdit, 'work_place')
        self.reference_from_work = self.findChild(QTextEdit, 'reference_from_work')
        self.phone_number = self.findChild(QLineEdit, 'phone_number')
        self.recruiting_commitee_decision = self.findChild(QTextEdit, 'recruiting_commitee_decision')

        self.button_add_change = self.findChild(QPushButton, 'button_add_change')
        self.rb_add = self.findChild(QRadioButton, 'rb_add')
        self.rb_change = self.findChild(QRadioButton, 'rb_change')
        self.flmc_id = self.findChild(QLabel, 'flmc_id')

        self.rb_add.setDisabled(True)
        self.rb_change.setDisabled(True)
        self.flmc_id.setText('ID: ' + self.id)

        if not self.is_full:
            self.button_add_change.clicked.connect(self.add_personal_data)
            self.button_add_change.setText('Добавить')
            self.setWindowTitle('Добавление основных данных военнообязанного')
            self.rb_add.setChecked(True)
        else:
            self.button_add_change.clicked.connect(self.change_personal_data)
            self.button_add_change.setText('Изменить')
            self.setWindowTitle('Изменение основных данных военнообязанного')
            self.rb_change.setChecked(True)
            data_row = self.get_row('personal_file.xlsx')
            self.passport_series.setText(str(data_row[1]))
            self.passport_id.setText(str(data_row[2]))
            self.birthday_place.setText(str(data_row[3]))
            self.living_place.setText(str(data_row[4]))
            self.family_status.setText(str(data_row[5]))
            self.education.setText(str(data_row[6]))
            self.work_place.setText(str(data_row[7]))
            self.reference_from_work.setText(str(data_row[8]))
            self.phone_number.setText(str(data_row[9]))
            self.recruiting_commitee_decision.setText(str(data_row[10]))

    def add_personal_data(self):
        table = pd.read_excel('bases/personal_file.xlsx')
        lfmc_id = self.id
        passport_series = self.passport_series.text()
        passport_id = self.passport_id.text()
        birthday_place = self.birthday_place.text()
        living_place = self.living_place.text()
        family_status = self.family_status.text()
        education = self.education.text()
        work_place = self.work_place.text()
        reference_from_work = self.reference_from_work.toPlainText()
        phone_number = self.phone_number.text()
        recruiting_commitee_decision = self.recruiting_commitee_decision.toPlainText()
        params = [lfmc_id, passport_series, passport_id,
                  birthday_place, living_place, family_status,
                  education, work_place, reference_from_work,
                  phone_number, recruiting_commitee_decision]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            dict1 = {'Lfmc_id': lfmc_id,
                     'Passport_series': passport_series,
                     'Passport_id': passport_id,
                     'Birthday_place': birthday_place,
                     'Living_place': living_place,
                     'Family_status': family_status,
                     'Education': education,
                     'Work_place': work_place,
                     'Reference_from_work': reference_from_work,
                     'Phone_number': phone_number,
                     'Recruiting_commitee_decision': recruiting_commitee_decision}
            table = pd.concat([table, pd.DataFrame([dict1])], ignore_index=True)
            table.to_excel('bases/personal_file.xlsx', index=False)
            self.show_message('Основные данные военнообязанного с идентификатором: ' + str(lfmc_id) + ' добавлены.')

    def change_personal_data(self):
        table = pd.read_excel('bases/personal_file.xlsx')
        lfmc_id = self.id
        passport_series = self.passport_series.text()
        passport_id = self.passport_id.text()
        birthday_place = self.birthday_place.text()
        living_place = self.living_place.text()
        family_status = self.family_status.text()
        education = self.education.text()
        work_place = self.work_place.text()
        reference_from_work = self.reference_from_work.toPlainText()
        phone_number = self.phone_number.text()
        recruiting_commitee_decision = self.recruiting_commitee_decision.toPlainText()
        params = [lfmc_id, passport_series, passport_id,
                  birthday_place, living_place, family_status,
                  education, work_place, reference_from_work,
                  phone_number, recruiting_commitee_decision]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            dict2 = {'Lfmc_id': lfmc_id,
                     'Passport_series': passport_series,
                     'Passport_id': passport_id,
                     'Birthday_place': birthday_place,
                     'Living_place': living_place,
                     'Family_status': family_status,
                     'Education': education,
                     'Work_place': work_place,
                     'Reference_from_work': reference_from_work,
                     'Phone_number': phone_number,
                     'Recruiting_commitee_decision': recruiting_commitee_decision}
            table.iloc[int(lfmc_id) - 1] = dict2
            table.to_excel('bases/personal_file.xlsx', index=False)
            self.show_message('Данные военнообязанного с ID: ' + str(lfmc_id) + ' изменены.')
            self.close()

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)

        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()


class ParentsInfo(QMainWindow):
    def __init__(self, incoming_id, is_full):
        super(ParentsInfo, self).__init__()
        self.id = incoming_id
        self.is_full = is_full
        uic.loadUi('parentsInfo.ui', self)

        self.father_full_name = self.findChild(QLineEdit, 'father_full_name')
        self.father_birthday = self.findChild(QLineEdit, 'father_birthday')
        self.father_birthday_place = self.findChild(QLineEdit, 'father_birthday_place')
        self.father_work_place = self.findChild(QLineEdit, 'father_work_place')
        self.mother_full_name = self.findChild(QLineEdit, 'mother_full_name')
        self.mother_birthday = self.findChild(QLineEdit, 'mother_birthday')
        self.mother_birthday_place = self.findChild(QLineEdit, 'mother_birthday_place')
        self.mother_work_place = self.findChild(QLineEdit, 'mother_work_place')

        self.button_add_change = self.findChild(QPushButton, 'button_add_change')
        self.rb_add = self.findChild(QRadioButton, 'rb_add')
        self.rb_change = self.findChild(QRadioButton, 'rb_change')
        self.flmc_id = self.findChild(QLabel, 'flmc_id')

        self.rb_add.setDisabled(True)
        self.rb_change.setDisabled(True)
        self.flmc_id.setText('ID: ' + self.id)

        if not self.is_full:
            self.button_add_change.clicked.connect(self.add_parents_info)
            self.button_add_change.setText('Добавить')
            self.setWindowTitle('Добавление данных о членах семьи военнообязанного')
            self.rb_add.setChecked(True)
        else:
            self.button_add_change.clicked.connect(self.change_parents_info)
            self.button_add_change.setText('Изменить')
            self.setWindowTitle('Изменение данных о членах семьи военнообязанного')
            self.rb_change.setChecked(True)
            data_row = self.get_row('parents_info.xlsx')
            self.father_full_name.setText(str(data_row[1]))
            self.father_birthday.setText(str(data_row[2]))
            self.father_birthday_place.setText(str(data_row[3]))
            self.father_work_place.setText(str(data_row[4]))
            self.mother_full_name.setText(str(data_row[5]))
            self.mother_birthday.setText(str(data_row[6]))
            self.mother_birthday_place.setText(str(data_row[7]))
            self.mother_work_place.setText(str(data_row[8]))

    def add_parents_info(self):
        table = pd.read_excel('bases/parents_info.xlsx')
        lfmc_id = self.id
        father_full_name = self.father_full_name.text()
        father_birthday = self.father_birthday.text()
        father_birthday_place = self.father_birthday_place.text()
        father_work_place = self.father_work_place.text()
        mother_full_name = self.mother_full_name.text()
        mother_birthday = self.mother_birthday.text()
        mother_birthday_place = self.mother_birthday_place.text()
        mother_work_place = self.mother_work_place.text()
        params = [lfmc_id, father_full_name, father_birthday,
                  father_birthday_place, father_work_place, mother_full_name,
                  mother_birthday, mother_birthday_place, mother_work_place]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            dict4 = {'Lfmc_id': lfmc_id,
                     'Father_full_name': father_full_name,
                     'Father_birthday': father_birthday,
                     'Father_birthday_place': father_birthday_place,
                     'Father_work_place': father_work_place,
                     'Mother_full_name': mother_full_name,
                     'Mother_birthday': mother_birthday,
                     'Mother_birthday_place': mother_birthday_place,
                     'Mother_work_place': mother_work_place}
            table = pd.concat([table, pd.DataFrame([dict4])], ignore_index=True)
            table.to_excel('bases/parents_info.xlsx', index=False)
            self.show_message('Данные о членах семьи военнообязанного с идентификатором: ' + str(lfmc_id) + ' добавлены.')
            self.close()

    def change_parents_info(self):
        table = pd.read_excel('bases/parents_info.xlsx')
        lfmc_id = self.id
        father_full_name = self.father_full_name.text()
        father_birthday = self.father_birthday.text()
        father_birthday_place = self.father_birthday_place.text()
        father_work_place = self.father_work_place.text()
        mother_full_name = self.mother_full_name.text()
        mother_birthday = self.mother_birthday.text()
        mother_birthday_place = self.mother_birthday_place.text()
        mother_work_place = self.mother_work_place.text()
        params = [lfmc_id, father_full_name, father_birthday,
                  father_birthday_place, father_work_place, mother_full_name,
                  mother_birthday, mother_birthday_place, mother_work_place]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue
        if no_empty_lines:
            index = table.index[table['Lfmc_id'] == int(lfmc_id)][0]
            table.at[index, 'Lfmc_id'] = lfmc_id
            table.at[index, 'Father_full_name'] = father_full_name
            table.at[index, 'Father_birthday'] = father_birthday
            table.at[index, 'Father_birthday_place'] = father_birthday_place
            table.at[index, 'Father_work_place'] = father_work_place
            table.at[index, 'Mother_full_name'] = mother_full_name
            table.at[index, 'Mother_birthday'] = mother_birthday
            table.at[index, 'Mother_birthday_place'] = mother_birthday_place
            table.at[index, 'Mother_work_place'] = mother_work_place
            table.to_excel('bases/parents_info.xlsx', index=False)
            self.show_message('Данные военнообязанного с ID: ' + str(lfmc_id) + ' изменены.')
            self.close()

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)

        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()


class MedicalDocument(QMainWindow):
    def __init__(self, incoming_id, is_full):
        super(MedicalDocument, self).__init__()
        self.id = incoming_id
        self.is_full = is_full
        uic.loadUi('medicalDocument.ui', self)

        self.medical_examination_date = self.findChild(QLineEdit, 'medical_examination_date')
        self.examination_place = self.findChild(QLineEdit, 'examination_place')
        self.complaints = self.findChild(QTextEdit, 'complaints')
        self.anamnesis = self.findChild(QTextEdit, 'anamnesis')
        self.objective_research_data = self.findChild(QTextEdit, 'objective_research_data')
        self.examination_results = self.findChild(QTextEdit, 'examination_results')
        self.diagnosis = self.findChild(QTextEdit, 'diagnosis')
        self.doctor_full_name = self.findChild(QLineEdit, 'doctor_full_name')
        self.medical_speciality = self.findChild(QLineEdit, 'medical_speciality')

        self.button_add_change = self.findChild(QPushButton, 'button_add_change')
        self.rb_add = self.findChild(QRadioButton, 'rb_add')
        self.rb_change = self.findChild(QRadioButton, 'rb_change')
        self.flmc_id = self.findChild(QLabel, 'flmc_id')

        self.rb_add.setDisabled(True)
        self.rb_change.setDisabled(True)
        self.flmc_id.setText('ID: ' + self.id)

        if not self.is_full:
            self.button_add_change.clicked.connect(self.add_med_doc)
            self.button_add_change.setText('Добавить')
            self.setWindowTitle('Добавление основных данных военнообязанного')
            self.rb_add.setChecked(True)
        else:
            self.button_add_change.clicked.connect(self.change_med_doc)
            self.button_add_change.setText('Изменить')
            self.setWindowTitle('Изменение основных данных военнообязанного')
            self.rb_change.setChecked(True)
            data_row = self.get_row('medical_documents.xlsx')
            self.medical_examination_date.setText(str(data_row[1]))
            self.examination_place.setText(str(data_row[2]))
            self.complaints.setText(str(data_row[3]))
            self.anamnesis.setText(str(data_row[4]))
            self.objective_research_data.setText(str(data_row[5]))
            self.examination_results.setText(str(data_row[6]))
            self.diagnosis.setText(str(data_row[7]))
            self.doctor_full_name.setText(str(data_row[8]))
            self.medical_speciality.setText(str(data_row[9]))

    def add_med_doc(self):
        table = pd.read_excel('bases/medical_documents.xlsx')
        lfmc_id = self.id
        medical_examination_date = self.medical_examination_date.text()
        examination_place = self.examination_place.text()
        complaints = self.complaints.toPlainText()
        anamnesis = self.anamnesis.toPlainText()
        objective_research_data = self.objective_research_data.toPlainText()
        examination_results = self.examination_results.toPlainText()
        diagnosis = self.diagnosis.toPlainText()
        doctor_full_name = self.doctor_full_name.text()
        medical_speciality = self.medical_speciality.text()
        params = [lfmc_id, medical_examination_date, examination_place,
                  complaints, anamnesis, objective_research_data,
                  examination_results, diagnosis, doctor_full_name,
                  medical_speciality]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            dict3 = {'Lfmc_id': lfmc_id,
                     'Medical_examination_date': medical_examination_date,
                     'Examination_place': examination_place,
                     'Complaints': complaints,
                     'Anamnesis': anamnesis,
                     'Objective_research_data': objective_research_data,
                     'Examination_results': examination_results,
                     'Diagnosis': diagnosis,
                     'Doctor_full_name': doctor_full_name,
                     'Medical_speciality': medical_speciality}
            table = pd.concat([table, pd.DataFrame([dict3])], ignore_index=True)
            table.to_excel('bases/medical_documents.xlsx', index=False)
            self.show_message('Медицинский документ военнообязанного с идентификатором: ' + str(lfmc_id) + ' добавлен.')
            self.close()

    def change_med_doc(self):
        table = pd.read_excel('bases/medical_documents.xlsx')
        lfmc_id = self.id
        medical_examination_date = self.medical_examination_date.text()
        examination_place = self.examination_place.text()
        complaints = self.complaints.toPlainText()
        anamnesis = self.anamnesis.toPlainText()
        objective_research_data = self.objective_research_data.toPlainText()
        examination_results = self.examination_results.toPlainText()
        diagnosis = self.diagnosis.toPlainText()
        doctor_full_name = self.doctor_full_name.text()
        medical_speciality = self.medical_speciality.text()
        params = [lfmc_id, medical_examination_date, examination_place,
                  complaints, anamnesis, objective_research_data,
                  examination_results, diagnosis, doctor_full_name,
                  medical_speciality]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue
        if no_empty_lines:
            index = table.index[table['Lfmc_id'] == int(lfmc_id)][0]
            table.at[index, 'Lfmc_id'] = lfmc_id
            table.at[index, 'Medical_examination_date'] = medical_examination_date
            table.at[index, 'Examination_place'] = examination_place
            table.at[index, 'Complaints'] = complaints
            table.at[index, 'Anamnesis'] = anamnesis
            table.at[index, 'Objective_research_data'] = objective_research_data
            table.at[index, 'Examination_results'] = examination_results
            table.at[index, 'Diagnosis'] = diagnosis
            table.at[index, 'Doctor_full_name'] = doctor_full_name
            table.at[index, 'Medical_speciality'] = medical_speciality
            table.to_excel('bases/medical_documents.xlsx', index=False)
            self.show_message('Данные военнообязанного с ID: ' + str(lfmc_id) + ' изменены.')
            self.close()

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)

        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()


class ArmyOrder(QMainWindow):
    def __init__(self, incoming_id, is_full):
        super(ArmyOrder, self).__init__()
        self.id = incoming_id
        self.is_full = is_full
        uic.loadUi('armyOrder.ui', self)

        self.visit_date = self.findChild(QLineEdit, 'visit_date')
        self.address = self.findChild(QLineEdit, 'address')
        self.visit_reason = self.findChild(QLineEdit, 'visit_reason')

        self.button_add_change = self.findChild(QPushButton, 'button_add_change')
        self.rb_add = self.findChild(QRadioButton, 'rb_add')
        self.rb_change = self.findChild(QRadioButton, 'rb_change')
        self.flmc_id = self.findChild(QLabel, 'flmc_id')

        self.rb_add.setDisabled(True)
        self.rb_change.setDisabled(True)
        self.flmc_id.setText('ID: ' + self.id)

        if not self.is_full:
            self.button_add_change.clicked.connect(self.add_army_order)
            self.button_add_change.setText('Добавить')
            self.setWindowTitle('Добавление повестки военнообязанному')
            self.rb_add.setChecked(True)
        else:
            self.button_add_change.clicked.connect(self.change_army_order)
            self.button_add_change.setText('Изменить')
            self.setWindowTitle('Изменение повестки военнообязанного')
            self.rb_change.setChecked(True)
            data_row = self.get_row('army_order.xlsx')
            self.visit_date.setText(str(data_row[1]))
            self.address.setText(str(data_row[2]))
            self.visit_reason.setText(str(data_row[3]))

    def add_army_order(self):
        table = pd.read_excel('bases/army_order.xlsx')
        lfmc_id = self.id
        visit_date = self.visit_date.text()
        address = self.address.text()
        visit_reason = self.visit_reason.text()
        params = [lfmc_id, visit_date, address, visit_reason]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            dict4 = {'Lfmc_id': lfmc_id,
                     'Visit_date': visit_date,
                     'Address': address,
                     'Visit_reason': visit_reason}
            table = pd.concat([table, pd.DataFrame([dict4])], ignore_index=True)
            table.to_excel('bases/army_order.xlsx', index=False)
            self.show_message('Повестка военнообязанного с идентификатором: ' + str(lfmc_id) + ' добавлена.')
            self.close()

    def change_army_order(self):
        table = pd.read_excel('bases/army_order.xlsx')
        lfmc_id = self.id
        visit_date = self.visit_date.text()
        address = self.address.text()
        visit_reason = self.visit_reason.text()
        params = [lfmc_id, visit_date, address, visit_reason]
        no_empty_lines = False
        for option in params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue
        if no_empty_lines:
            index = table.index[table['Lfmc_id'] == int(lfmc_id)][0]
            table.at[index, 'Lfmc_id'] = lfmc_id
            table.at[index, 'Visit_date'] = visit_date
            table.at[index, 'Address'] = address
            table.at[index, 'Visit_reason'] = visit_reason
            table.to_excel('bases/army_order.xlsx', index=False)
            self.show_message('Данные военнообязанного с ID: ' + str(lfmc_id) + ' изменены.')
            self.close()

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)

        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()


class PersonalFile(QMainWindow):
    def __init__(self, incoming_id):
        super(PersonalFile, self).__init__()
        self.id = incoming_id
        uic.loadUi('personalFile.ui', self)

        self.flmc_id = self.findChild(QLabel, 'flmc_id')
        self.flmc_id.setText('ID: ' + self.id)

        self.act_create = self.findChild(QAction, 'create_personal_file')
        self.act_add_change_personal_file = self.findChild(QAction, 'add_change_personal_file')
        self.act_add_change_parents_info = self.findChild(QAction, 'add_change_parents_info')
        self.act_add_change_medical_document = self.findChild(QAction, 'add_change_medical_document')
        self.act_add_change_army_order = self.findChild(QAction, 'add_change_army_order')

        self.act_create.triggered.connect(self.view_info)
        self.act_add_change_personal_file.triggered.connect(self.addch_personal_file)
        self.act_add_change_parents_info.triggered.connect(self.addch_parents_info)
        self.act_add_change_medical_document.triggered.connect(self.addch_medical_document)
        self.act_add_change_army_order.triggered.connect(self.addch_army_order)

        self.full_name = self.findChild(QLabel, 'full_name')
        self.birthday_date = self.findChild(QLabel, 'birthday_date')
        self.health_category = self.findChild(QLabel, 'health_category')
        self.military_specialty = self.findChild(QLabel, 'military_specialty')
        self.combat_experience = self.findChild(QLabel, 'combat_experience')

        self.passport_series = self.findChild(QLabel, 'passport_series')
        self.passport_id = self.findChild(QLabel, 'passport_id')
        self.birthday_place = self.findChild(QLabel, 'birthday_place')
        self.living_place = self.findChild(QLabel, 'living_place')
        self.family_status = self.findChild(QLabel, 'family_status')
        self.education = self.findChild(QLabel, 'education')
        self.work_place = self.findChild(QLabel, 'work_place')
        self.reference_from_work = self.findChild(QLabel, 'reference_from_work')
        self.phone_number = self.findChild(QLabel, 'phone_number')
        self.commitee_decision = self.findChild(QLabel, 'commitee_decision')

        self.father_full_name = self.findChild(QLabel, 'father_full_name')
        self.father_birthday_date = self.findChild(QLabel, 'father_birthday_date')
        self.father_birthday_place = self.findChild(QLabel, 'father_birthday_place')
        self.father_work_place = self.findChild(QLabel, 'father_work_place')
        self.mother_full_name = self.findChild(QLabel, 'mother_full_name')
        self.mother_birthday_date = self.findChild(QLabel, 'mother_birthday_date')
        self.mother_birthday_place = self.findChild(QLabel, 'mother_birthday_place')
        self.mother_work_place = self.findChild(QLabel, 'mother_work_place')

        self.visit_date = self.findChild(QLabel, 'visit_date')
        self.comissariat_address = self.findChild(QLabel, 'comissariat_address')
        self.visit_reason = self.findChild(QLabel, 'visit_reason')

        self.examination_date = self.findChild(QLabel, 'examination_date')
        self.examination_place = self.findChild(QLabel, 'examination_place')
        self.doctor_full_name = self.findChild(QLabel, 'doctor_full_name')
        self.doctor_specialty = self.findChild(QLabel, 'doctor_specialty')
        self.complaints = self.findChild(QLabel, 'complaints')
        self.anamnesis = self.findChild(QLabel, 'anamnesis')
        self.objective_research_data = self.findChild(QLabel, 'objective_research_data')
        self.examination_results = self.findChild(QLabel, 'examination_results')
        self.diagnosis = self.findChild(QLabel, 'diagnosis')

        self.is_created = False
        self.pers_data = PersonalData(self.id, False)
        self.parents_info = ParentsInfo(self.id, False)
        self.med_doc = MedicalDocument(self.id, False)
        self.arm_order = ArmyOrder(self.id, False)

    def view_info(self):
        self.is_created = True
        lfmc_row = self.get_row('lfmc.xlsx')
        self.full_name.setText(lfmc_row[1] + ' ' + lfmc_row[2] + ' ' + lfmc_row[3])
        self.birthday_date.setText('Дата рождения: ' + str(lfmc_row[4])[0:10])
        self.health_category.setText('Категория годности: ' + lfmc_row[5])
        self.military_specialty.setText('Военная специальность: ' + lfmc_row[6])
        self.combat_experience.setText('Военный опыт: ' + lfmc_row[7])

        is_full = True

        if self.is_flms_data_exist('personal_file.xlsx'):
            personal_file_row = self.get_row('personal_file.xlsx')
            self.passport_series.setText('Серия: ' + str(personal_file_row[1]))
            self.passport_id.setText('Номер: ' + str(personal_file_row[2]))
            self.birthday_place.setText('Место рождения: ' + str(personal_file_row[3]))
            self.living_place.setText('Место проживания: ' + str(personal_file_row[4]))
            self.family_status.setText('Семейное положение: ' + str(personal_file_row[5]))
            self.education.setText('Образование: ' + str(personal_file_row[6]))
            self.work_place.setText('Место работы: ' + str(personal_file_row[7]))
            self.reference_from_work.setText('Характеристика с работы: ' + str(personal_file_row[8]))
            self.phone_number.setText('Номер телефона: ' + str(personal_file_row[9]))
            self.commitee_decision.setText('Решение комиссии: ' + str(personal_file_row[10]))
            self.pers_data = PersonalData(self.id, True)
        else:
            is_full = False
            self.pers_data = PersonalData(self.id, False)

        if self.is_flms_data_exist('parents_info.xlsx'):
            parents_info_row = self.get_row('parents_info.xlsx')
            self.father_full_name.setText(str(parents_info_row[1]))
            self.father_birthday_date.setText('Дата рождения: ' + str(parents_info_row[2])[0:10])
            self.father_birthday_place.setText('Место рождения: ' + str(parents_info_row[3]))
            self.father_work_place.setText('Место работы: ' + str(parents_info_row[4]))
            self.mother_full_name.setText(str(parents_info_row[5]))
            self.mother_birthday_date.setText('Дата рождения: ' + str(parents_info_row[6])[0:10])
            self.mother_birthday_place.setText('Место рождения: ' + str(parents_info_row[7]))
            self.mother_work_place.setText('Место работы: ' + str(parents_info_row[8]))
            self.parents_info = ParentsInfo(self.id, True)
        else:
            is_full = False
            self.parents_info = ParentsInfo(self.id, False)

        if self.is_flms_data_exist('army_order.xlsx'):
            army_order_row = self.get_row('army_order.xlsx')
            self.visit_date.setText('Дата посещения: ' + str(army_order_row[1]))
            self.comissariat_address.setText('Адрес военкомата: ' + str(army_order_row[2]))
            self.visit_reason.setText('Причина посещения: ' + str(army_order_row[3]))
            self.arm_order = ArmyOrder(self.id, True)
        else:
            is_full = False
            self.arm_order = ArmyOrder(self.id, False)

        if self.is_flms_data_exist('medical_documents.xlsx'):
            medical_documents_row = self.get_row('medical_documents.xlsx')
            self.examination_date.setText('Дата проведения осмотра: ' + str(medical_documents_row[1])[0:10])
            self.examination_place.setText('Место проведения осмотра: ' + str(medical_documents_row[2]))
            self.doctor_full_name.setText('ФИО врача: ' + str(medical_documents_row[8]))
            self.doctor_specialty.setText('Специальность врача: ' + str(medical_documents_row[9]))
            self.complaints.setText('Жалобы: ' + str(medical_documents_row[3]))
            self.anamnesis.setText('Анамнез: ' + str(medical_documents_row[4]))
            self.objective_research_data.setText('Данные объективного исследования: ' + str(medical_documents_row[5]))
            self.examination_results.setText('Результаты осмотра: ' + str(medical_documents_row[6]))
            self.diagnosis.setText('Диагноз: ' + str(medical_documents_row[7]))
            self.med_doc = MedicalDocument(self.id, True)
        else:
            is_full = False
            self.med_doc = MedicalDocument(self.id, False)

        if not is_full:
            self.show_message('Для данного военнообязанного найдены не все данные, дополните личное дело.')

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)
        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def is_flms_data_exist(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)
        is_exist = False
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                is_exist = True
        return is_exist

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()

    def addch_personal_file(self):
        if self.is_created:
            self.pers_data.show()
        else:
            self.show_error('Личное дело не сформировано')

    def addch_parents_info(self):
        if self.is_created:
            self.parents_info.show()
        else:
            self.show_error('Личное дело не сформировано')

    def addch_medical_document(self):
        if self.is_created:
            self.med_doc.show()
        else:
            self.show_error('Личное дело не сформировано')

    def addch_army_order(self):
        if self.is_created:
            self.arm_order.show()
        else:
            self.show_error('Личное дело не сформировано')

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)


class AddChangeFLMC(QMainWindow):
    def __init__(self, incoming_id):
        super(AddChangeFLMC, self).__init__()
        self.id = incoming_id
        uic.loadUi('addChangeFlmc.ui', self)

        self.lfmc_surname = self.findChild(QLineEdit, 'flmc_surname')
        self.lfmc_name = self.findChild(QLineEdit, 'flmc_name')
        self.lfmc_patronymic = self.findChild(QLineEdit, 'flmc_patronymic')
        self.lfmc_birthday_date = self.findChild(QLineEdit, 'flmc_birthday_date')
        self.lfmc_health_category = self.findChild(QLineEdit, 'flmc_health_category')
        self.lfmc_military_speciality = self.findChild(QLineEdit, 'flmc_military_speciality')
        self.lfmc_combat_experience = self.findChild(QLineEdit, 'flmc_combat_experience')

        self.button_add_change = self.findChild(QPushButton, 'button_add_change')
        self.rb_add = self.findChild(QRadioButton, 'rb_add')
        self.rb_change = self.findChild(QRadioButton, 'rb_change')
        self.flmc_id = self.findChild(QLabel, 'flmc_id')

        self.rb_add.setDisabled(True)
        self.rb_change.setDisabled(True)
        self.flmc_id.setText('ID: ' + self.id)
        if self.id == '—':
            self.button_add_change.clicked.connect(self.add_flmc)
            self.button_add_change.setText('Добавить')
            self.setWindowTitle('Добавление данных военнообязанного')
            self.rb_add.setChecked(True)
        else:
            self.button_add_change.clicked.connect(self.change_flmc)
            self.button_add_change.setText('Изменить')
            self.setWindowTitle('Изменение данных военнообязанного')
            self.rb_change.setChecked(True)
            flmc_row = self.get_row('lfmc.xlsx')
            self.lfmc_surname.setText(str(flmc_row[1]))
            self.lfmc_name.setText(str(flmc_row[2]))
            self.lfmc_patronymic.setText(str(flmc_row[3]))
            self.lfmc_birthday_date.setText(str(flmc_row[4]))
            self.lfmc_health_category.setText(str(flmc_row[5]))
            self.lfmc_military_speciality.setText(str(flmc_row[6]))
            self.lfmc_combat_experience.setText(str(flmc_row[7]))

    def add_flmc(self):
        lfmc_table = pd.read_excel('bases/lfmc.xlsx')
        lfmc_id = lfmc_table.shape[0] + 1
        lfmc_surname = self.lfmc_surname.text()
        lfmc_name = self.lfmc_name.text()
        lfmc_patronymic = self.lfmc_patronymic.text()
        lfmc_birthday_date = self.lfmc_birthday_date.text()
        lfmc_health_category = self.lfmc_health_category.text()
        lfmc_military_speciality = self.lfmc_military_speciality.text()
        lfmc_combat_experience = self.lfmc_combat_experience.text()
        lfmc_params = [lfmc_id, lfmc_surname, lfmc_name,
                       lfmc_patronymic, lfmc_birthday_date, lfmc_health_category,
                       lfmc_military_speciality, lfmc_combat_experience]
        no_empty_lines = False
        for option in lfmc_params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            lfmc_dict = {'Lfmc_id': lfmc_id,
                         'Surname': lfmc_surname,
                         'Name': lfmc_name,
                         'Patronymic': lfmc_patronymic,
                         'Birthday_date': lfmc_birthday_date,
                         'Health_category': lfmc_health_category,
                         'Military_speciality': lfmc_military_speciality,
                         'Combat_experience': lfmc_combat_experience}
            lfmc_table = pd.concat([lfmc_table, pd.DataFrame([lfmc_dict])], ignore_index=True)
            lfmc_table.to_excel('bases/lfmc.xlsx', index=False)
            self.show_message('Военнообязанный добавлен, его идентификатор: ' + str(lfmc_id) + '.')
            self.close()

    def change_flmc(self):
        lfmc_table = pd.read_excel('bases/lfmc.xlsx')
        lfmc_id = self.id
        lfmc_surname = self.lfmc_surname.text()
        lfmc_name = self.lfmc_name.text()
        lfmc_patronymic = self.lfmc_patronymic.text()
        lfmc_birthday_date = self.lfmc_birthday_date.text()
        lfmc_health_category = self.lfmc_health_category.text()
        lfmc_military_speciality = self.lfmc_military_speciality.text()
        lfmc_combat_experience = self.lfmc_combat_experience.text()
        lfmc_params = [lfmc_id, lfmc_surname, lfmc_name,
                       lfmc_patronymic, lfmc_birthday_date, lfmc_health_category,
                       lfmc_military_speciality, lfmc_combat_experience]
        no_empty_lines = False
        for option in lfmc_params:
            if option == '':
                self.show_error('Не все поля заполнены.')
                no_empty_lines = False
                break
            else:
                no_empty_lines = True
                continue

        if no_empty_lines:
            index = lfmc_table.index[lfmc_table['Lfmc_id'] == int(lfmc_id)][0]
            lfmc_table.at[index, 'Lfmc_id'] = lfmc_id
            lfmc_table.at[index, 'Surname'] = lfmc_surname
            lfmc_table.at[index, 'Name'] = lfmc_name
            lfmc_table.at[index, 'Patronymic'] = lfmc_patronymic
            lfmc_table.at[index, 'Birthday_date'] = lfmc_birthday_date
            lfmc_table.at[index, 'Health_category'] = lfmc_health_category
            lfmc_table.at[index, 'Military_speciality'] = lfmc_military_speciality
            lfmc_table.at[index, 'Combat_experience'] = lfmc_combat_experience
            lfmc_table.to_excel('bases/lfmc.xlsx', index=False)
            self.show_message('Данные военнообязанного с ID: ' + str(lfmc_id) + ' изменены.')
            self.close()

    def get_row(self, file_name):
        data_frame = pd.read_excel('bases/' + file_name)
        data_row = ''
        for row in range(0, data_frame.shape[0]):
            if str(data_frame.iloc[row][0]) == str(self.id):
                data_row = data_frame.iloc[row]
        return data_row

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)

    def show_message(self, message):
        message_box = QMessageBox()
        message_box.setIcon(QMessageBox.Information)
        message_box.setText(message)
        message_box.setWindowTitle('Уведомление')
        message_box.setStandardButtons(QMessageBox.Ok)
        message_box.exec()


class AboutProgram(QMainWindow):
    def __init__(self):
        super(AboutProgram, self).__init__()
        uic.loadUi('aboutProgram.ui', self)


class Help(QMainWindow):
    def __init__(self):
        super(Help, self).__init__()
        uic.loadUi('helpUI.ui', self)


class MainTable(QMainWindow):
    def __init__(self):
        super(MainTable, self).__init__()
        uic.loadUi('mainTable.ui', self)

        self.mainTableWidget = self.findChild(QTableWidget, 'mainTableWidget')
        self.action_help = self.findChild(QAction, 'help')
        self.action_about_program = self.findChild(QAction, 'about_program')
        self.action_load = self.findChild(QAction, 'load_base')
        self.action_update = self.findChild(QAction, 'update_base')
        self.action_report = self.findChild(QAction, 'view_report')
        self.action_add_change_flmc = self.findChild(QAction, 'add_change_flmc')
        self.lfmc_id = self.findChild(QLabel, 'lfmc_id')
        self.reset_selection = self.findChild(QPushButton, 'reset_selection')
        self.prompt = self.findChild(QLabel, 'prompt')

        self.action_help.triggered.connect(self.act_help)
        self.action_about_program.triggered.connect(self.act_about_program)
        self.action_load.triggered.connect(self.act_load)
        self.action_update.triggered.connect(self.act_update)
        self.action_report.triggered.connect(self.act_report)
        self.action_add_change_flmc.triggered.connect(self.act_add_flmc)
        self.mainTableWidget.itemSelectionChanged.connect(self.cell_clicked)
        self.reset_selection.clicked.connect(self.reset)

        self.prompt.hide()
        self.selected_id = '—'
        self.aboutProgram = AboutProgram()
        self.help = Help()
        self.report = PersonalFile(self.selected_id)
        self.add_change = AddChangeFLMC(self.selected_id)
        self.is_loaded = False

    def reset(self):
        self.selected_id = '—'
        self.lfmc_id.setText('Выделен военнообязанный с идентификатором: ' + self.selected_id)

    def act_add_flmc(self):
        if self.is_loaded:
            self.add_change = AddChangeFLMC(self.selected_id)
            self.prompt.show()
            self.add_change.show()
        else:
            self.show_error('База данных не загружена.')

    def act_report(self):
        if self.mainTableWidget.rowCount() != 0:
            if self.selected_id != '—':
                self.report.show()
            else:
                self.show_error('Военнообязанный не выбран.')
        else:
            self.show_error('База данных не загружена.')

    def act_help(self):
        self.help.show()

    def act_about_program(self):
        self.aboutProgram.show()

    def act_load(self):
        if not self.is_loaded:
            lfmc_table = pd.read_excel('bases/lfmc.xlsx')
            self.mainTableWidget.setColumnCount(len(lfmc_table.columns))
            self.mainTableWidget.setRowCount(lfmc_table.shape[0])
            column_header = ("Идентификатор", "Фамилия", "Имя", "Отчество", "Дата рождения", "Категория годности",
                             "Военная специальность", "Боевой опыт")
            self.mainTableWidget.setHorizontalHeaderLabels(column_header)
            self.mainTableWidget.verticalHeader().setVisible(False)
            self.mainTableWidget.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
            self.is_loaded = True
            self.act_update()
        else:
            self.show_error('База данных уже загружена, используйте "Обновить базу".')

    def act_update(self):
        if self.is_loaded:
            self.prompt.hide()
            lfmc_table = pd.read_excel('bases/lfmc.xlsx', usecols=['Lfmc_id', 'Surname', 'Name', 'Patronymic',
                                                                   'Birthday_date', 'Health_category',
                                                                   'Military_speciality', 'Combat_experience'])
            self.mainTableWidget.setColumnCount(len(lfmc_table.columns))
            self.mainTableWidget.setRowCount(lfmc_table.shape[0])
            for row in range(0, lfmc_table.shape[0]):
                column = 0
                for col_name, data in lfmc_table.items():
                    item = QTableWidgetItem(str(data[row]))
                    item.setTextAlignment(QtCore.Qt.AlignCenter)
                    self.mainTableWidget.setItem(row, column, item)
                    column += 1
        else:
            self.show_error('База данных не загружена.')

    def cell_clicked(self):
        self.selected_id = self.mainTableWidget.model().index(self.mainTableWidget.currentRow(), 0).data()
        self.report = PersonalFile(self.selected_id)
        self.lfmc_id.setText('Выделен военнообязанный с идентификатором: ' + self.selected_id)

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)


class UI(QMainWindow):
    def __init__(self):
        super(UI, self).__init__()
        uic.loadUi("LoginForm.ui", self)
        self.show()

        self.loginForm_button = self.findChild(QPushButton, 'loginForm_button')
        self.loginForm_login = self.findChild(QLineEdit, 'loginForm_login')
        self.loginForm_password = self.findChild(QLineEdit, 'loginForm_password')

        self.loginForm_button.clicked.connect(self.log_in)

        self.mainTable = MainTable()

    def log_in(self):
        logins_data = pd.read_excel('bases/registration_data.xlsx', usecols=['Login'])
        passwords_data = pd.read_excel('bases/registration_data.xlsx', usecols=['Password'])
        login = self.loginForm_login.text()
        password = self.loginForm_password.text()
        for i in range(0, logins_data.shape[0]):
            if logins_data.iloc[i][0] == login:
                if passwords_data.iloc[i][0] == password:
                    self.mainTable.show()
                    self.close()
                    break
                else:
                    self.show_error("Неверный логин или пароль.")
                    break
            if i == logins_data.shape[0] - 1:
                self.show_error("Неверный логин или пароль.")

    def show_error(self, message):
        QMessageBox().critical(self, 'Ошибка', message, QMessageBox.Ok)


if __name__ == '__main__':
    app = QApplication(sys.argv)
    UIWindow = UI()
    app.exec_()
